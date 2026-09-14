"""Generate a leadership briefing from the live local ClickHouse warehouse."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from africa_pulse.settings import Settings
from africa_pulse.warehouse.client import get_client

OUT = Path(__file__).with_name("Africa_Pulse_Live_Platform_Briefing.pptx")
NAVY = RGBColor(16, 42, 67)
TEAL = RGBColor(11, 110, 153)
GREEN = RGBColor(19, 138, 114)
AMBER = RGBColor(180, 83, 9)
RED = RGBColor(180, 35, 24)
INK = RGBColor(30, 54, 78)
MUTED = RGBColor(72, 101, 129)
LINE = RGBColor(217, 226, 236)
PALE = RGBColor(248, 251, 253)
WHITE = RGBColor(255, 255, 255)


def query(client, sql):
    return list(client.query(sql).named_results())


def add_text(slide, text, x, y, w, h, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    paragraph.space_after = Pt(0)
    return shape


def box(slide, x, y, w, h, fill=WHITE, line=LINE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def new_slide(prs, title, subtitle, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_text(slide, "AFRICA PULSE", 0.55, 0.26, 2.2, 0.25, 10, TEAL, True)
    add_text(slide, title, 0.55, 0.60, 12.0, 0.45, 26, NAVY, True)
    add_text(slide, subtitle, 0.55, 1.12, 12.0, 0.26, 11, MUTED)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.54), Inches(12.2), Inches(0.02))
    rule.fill.solid()
    rule.fill.fore_color.rgb = LINE
    rule.line.fill.background()
    add_text(slide, f"{page:02d}", 12.35, 7.05, 0.3, 0.16, 9, MUTED, True, PP_ALIGN.RIGHT)
    return slide


def metric_card(slide, title, value, detail, x, accent):
    box(slide, x, 2.10, 3.72, 2.10, PALE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.10), Inches(0.10), Inches(2.10))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, title, x + 0.27, 2.40, 3.10, 0.25, 13, INK, True)
    add_text(slide, value, x + 0.27, 2.90, 3.10, 0.43, 27, accent, True)
    add_text(slide, detail, x + 0.27, 3.52, 3.10, 0.35, 10, MUTED)


def main():
    client = get_client(Settings.from_environment())
    mobility = query(client, "SELECT city_id, median_congestion_ratio, valid_sample_coverage_pct FROM mart.city_mobility_daily FINAL ORDER BY local_date DESC LIMIT 1 BY city_id")
    scores = query(client, "SELECT city_id, weighted_coverage_pct, score_status FROM mart.city_intelligence_daily FINAL ORDER BY local_date DESC LIMIT 1 BY city_id")
    runs = query(client, "SELECT status, count() AS count FROM control.ingestion_run FINAL GROUP BY status")
    traffic_fact_count = query(client, "SELECT count() AS count FROM warehouse.fact_traffic_flow_observation FINAL")[0]["count"]
    latest_commercial = query(client, "SELECT status, error_message FROM control.ingestion_run FINAL WHERE source_id = 'openstreetmap_overpass_commercial_v1' ORDER BY started_at DESC LIMIT 1")[0]
    run_counts = {row["status"]: row["count"] for row in runs}
    city_names = {"lagos_ng": "Lagos", "abuja_ng": "Abuja", "cape_town_za": "Cape Town"}

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Africa Pulse Live Platform Briefing"

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_text(slide, "AFRICA PULSE", 0.7, 0.70, 4, 0.3, 14, RGBColor(119, 224, 208), True)
    add_text(slide, "Trusted city intelligence\nwith live operational evidence", 0.7, 1.35, 8.6, 1.3, 33, WHITE, True)
    add_text(slide, "Lagos, Abuja, and Cape Town | ClickHouse analytics platform", 0.7, 3.00, 7.7, 0.35, 16, RGBColor(211, 225, 233))
    box(slide, 9.35, 1.20, 2.95, 4.75, RGBColor(25, 57, 85), RGBColor(48, 93, 123))
    add_text(slide, "The decision standard", 9.75, 1.65, 2.15, 0.3, 12, RGBColor(119, 224, 208), True, PP_ALIGN.CENTER)
    add_text(slide, "Can we trust\nthe number?", 9.70, 2.35, 2.25, 0.8, 23, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, "Every metric has source, time, quality, and recovery evidence.", 9.70, 3.65, 2.25, 0.7, 13, RGBColor(211, 225, 233), False, PP_ALIGN.CENTER)
    add_text(slide, "Technical leadership briefing", 0.7, 7.04, 3.3, 0.18, 9, RGBColor(171, 194, 210))

    slide = new_slide(prs, "The platform is working end to end", "Live sources flow into ClickHouse facts, quality controls, and business marts.", 2)
    metric_card(slide, "Cities served", "3", "Lagos, Abuja, and Cape Town", 0.55, TEAL)
    metric_card(slide, "Traffic facts stored", str(traffic_fact_count), "Five configured road samples per city", 4.80, GREEN)
    metric_card(slide, "Operational source runs", str(sum(run_counts.values())), f"{run_counts.get('completed', 0)} completed; failures retained as evidence", 9.05, AMBER)
    add_text(slide, "What this proves", 0.55, 4.85, 2.0, 0.25, 14, NAVY, True)
    add_text(slide, "The implementation is not a mock-up. Real TomTom, Open-Meteo, FX, OpenStreetMap, and World Bank responses have been acquired, validated, and made queryable.", 0.55, 5.28, 11.7, 0.55, 17, INK)

    slide = new_slide(prs, "One comparable mobility source supports all three cities", "TomTom Traffic Flow provides a road-congestion proxy with identical fields and collection rules.", 3)
    for index, row in enumerate(mobility):
        name = city_names[row["city_id"]]
        congestion = float(row["median_congestion_ratio"]) * 100
        coverage = float(row["valid_sample_coverage_pct"])
        metric_card(slide, name, f"{congestion:.1f}%", f"Median congestion | {coverage:.0f}% validated sample coverage", 0.55 + index * 4.25, [TEAL, GREEN, AMBER][index])
    add_text(slide, "Decision: report this as a sampled road-congestion proxy, never as ridership or passenger demand.", 0.55, 5.02, 11.6, 0.3, 15, NAVY, True)
    add_text(slide, "Why: it is technically comparable across the portfolio, while its published coverage and confidence prevent over-interpretation.", 0.55, 5.52, 11.6, 0.35, 15, MUTED)

    slide = new_slide(prs, "The architecture protects evidence before it produces metrics", "Every layer has a single responsibility and contributes to traceability.", 4)
    stages = [("Acquire", "Source connectors", TEAL), ("Preserve", "Raw evidence + checksums", GREEN), ("Validate", "Quality rules + quarantine", AMBER), ("Integrate", "ClickHouse facts", RED), ("Serve", "Daily marts + dashboard", TEAL)]
    for index, (name, detail, color) in enumerate(stages):
        x = 0.55 + index * 2.53
        box(slide, x, 2.35, 2.05, 1.25, PALE)
        add_text(slide, name, x + 0.18, 2.65, 1.68, 0.22, 14, color, True, PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.18, 3.08, 1.68, 0.25, 10, MUTED, False, PP_ALIGN.CENTER)
        if index < 4:
            add_text(slide, ">", x + 2.10, 2.78, 0.3, 0.3, 20, TEAL, True, PP_ALIGN.CENTER)
    box(slide, 0.55, 4.60, 12.0, 1.05, RGBColor(235, 247, 245), RGBColor(193, 229, 221))
    add_text(slide, "Trust controls", 0.85, 4.93, 1.7, 0.2, 13, NAVY, True)
    add_text(slide, "Deterministic observation keys | run IDs | source checksums | freshness | validation outcomes | recovery history", 2.55, 4.93, 9.4, 0.22, 14, INK, True)

    slide = new_slide(prs, "Failure is recorded and recovery is demonstrable", "The platform retains operational history instead of hiding source problems.", 5)
    metric_card(slide, "Latest commercial run", latest_commercial["status"].upper(), "Rate-limited source failure is retained and visible", 0.55, RED)
    metric_card(slide, "Recovery design", "Buffered write", "Future snapshots publish only after all cities are acquired", 4.80, GREEN)
    metric_card(slide, "Duplicate rerun", "0 inserts", "Second TomTom run kept the fact count stable", 9.05, TEAL)
    add_text(slide, "Why this matters", 0.55, 4.85, 1.9, 0.25, 14, NAVY, True)
    add_text(slide, "An independent engineer can see what failed, inspect the source evidence, wait for the provider limit to clear, rerun the workflow, and prove that recovery did not duplicate records.", 0.55, 5.28, 11.7, 0.55, 17, INK)

    slide = new_slide(prs, "The City Intelligence Score refuses to overstate certainty", "Components remain visible, but a rank is withheld until the evidence is sufficient.", 6)
    for index, row in enumerate(scores):
        metric_card(slide, city_names[row["city_id"]], row["score_status"], f"{float(row['weighted_coverage_pct']):.0f}% weighted component coverage", 0.55 + index * 4.25, [TEAL, GREEN, AMBER][index])
    add_text(slide, "Current rule", 0.55, 4.90, 1.5, 0.24, 14, NAVY, True)
    add_text(slide, "A numeric score requires approved commercial-density denominators, seven days of FX history, and 28 days of trend history. Until then, `unavailable` is the correct decision-grade result.", 0.55, 5.32, 11.7, 0.52, 16, INK)

    slide = new_slide(prs, "The next decision is operational maturity, not new features", "The platform is demonstrable today; scheduled collection will turn the initial vertical slice into a historical decision asset.", 7)
    actions = [("1", "Schedule every two hours", "Build the FX and traffic history window", TEAL), ("2", "Monitor operational health", "Act on freshness, failure, and quality evidence", GREEN), ("3", "Approve score inputs", "Confirm municipal boundaries and commercial-density denominator", AMBER), ("4", "Scale after measurement", "Add workers or aggregates only when the workload proves the need", RED)]
    for index, (number, title, detail, color) in enumerate(actions):
        x = 0.65 + (index % 2) * 6.15
        y = 2.10 + (index // 2) * 1.62
        box(slide, x, y, 5.55, 1.15, PALE)
        add_text(slide, number, x + 0.25, y + 0.33, 0.35, 0.28, 18, color, True)
        add_text(slide, title, x + 0.80, y + 0.26, 4.35, 0.25, 14, NAVY, True)
        add_text(slide, detail, x + 0.80, y + 0.63, 4.35, 0.22, 11, MUTED)
    add_text(slide, "Decision requested: approve the operating cadence and the data-governance inputs needed for a fully published City Intelligence Score.", 0.65, 5.88, 11.65, 0.32, 15, NAVY, True, PP_ALIGN.CENTER)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
